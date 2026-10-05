"""
Build TemanJourney_Model_2027_v14.xlsx from v13 (v13 is never changed).

Run from financial-estimations/:   python analysis/build_v14.py
What v14 adds, all as live formulas:
  1. Timing (Inputs section 11): "pipeline first" - no new B2B or B2C events before the first delivery month
     (default Mar-27). Existing EAP and counseling clients continue.
  2. Growth plan (Inputs section 12): fair pay for full-time staff, extra IG ads, a full-time psychologist that starts
     with EAP contract 2, two new EAP contracts, a sales co-founder, and bank training days (priced as bank events).
  3. Outlook sheet: sales, total cost, net profit and cash each month, with one combined chart.
openpyxl drops cached values, so recalculate in Excel afterwards (analysis/excel_runner.py does this).
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as CL

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = ROOT / "TemanJourney_Model_2027_v13.xlsx", ROOT / "TemanJourney_Model_2027_v14.xlsx"
wb = openpyxl.load_workbook(SRC)
I, V, C, P, CF, S, T = (wb[n] for n in ["Inputs", "Volume", "Capacity_OPEX", "PnL", "CashFlow", "Summary", "To_Fill"])
COLS = range(5, 20)                     # E..S = Oct-26 .. Dec-27


def style_like(dst, src):
    dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def month(cell):                       # month index of an Inputs month label, 99 if the label is not found
    return f"IFERROR(MATCH(Inputs!${cell[0]}${cell[1:]},Volume!$E$4:$S$4,0),99)"


# ------------------------------------------------------------------ 1. Inputs sections 11 and 12
rows = [
    ("11. Timing - when new sales really start (v14)", None),
    ("pipe_on", "Pipeline-first timing? (1 = no new sales before the month below)", 1, "switch", "Your plan - confirm",
     "Oct 2026 to Feb 2027 is used to pitch and sign. Existing EAP and counseling clients continue."),
    ("sales_from", "First month new contracts are delivered", "Mar-27", "month", "Your plan - confirm",
     "Must match a month label (e.g. Mar-27). Earlier months get no new B2B events."),
    ("pipe_b2c", "Also stop B2C mass events before that month? (1 = yes)", 1, "switch", "Your plan - confirm",
     "1 = the team focuses on pitching. 0 = keep public events running."),
    ("12. Growth plan - fair pay and the hires in the growth deck (v14)", None),
    ("gp_on", "Growth plan switch (1 = on, 0 = v13 cost base)", 1, "switch", "Your plan - confirm",
     "Turns on every line in this section."),
    ("fair_on", "Pay full-time staff fairly? (1 = yes)", 1, "switch", "Your plan - confirm",
     "Minimum wage + BPJS + THR. Required by law for employees."),
    ("fair_start", "Fair pay start month", "Jan-27", "month", "Your plan - confirm", "Must match a month label."),
    ("oncost", "BPJS + THR on top of salary", 0.18, "%", "Estimate", "Employer BPJS + holiday bonus; confirm with a payroll provider."),
    ("fp_admin", "Fair gross pay: admin", 5.73, "Rp m / month", "Known", "Jakarta minimum wage 2026: Rp5,729,876."),
    ("fp_socmed", "Fair gross pay: social media specialist", 5.73, "Rp m / month", "Known", "Jakarta minimum wage 2026."),
    ("fp_lead", "Fair gross pay: lead marketing", 7.0, "Rp m / month", "Estimate", "Minimum wage + about 22% for a lead role."),
    ("fp_psych", "Fair gross pay: full-time psychologist (replaces PIC part-time)", 9.0, "Rp m / month", "Estimate",
     "Low end of Rp7m to Rp13m for licensed psychologists (JobStreet)."),
    ("ads_extra", "Extra IG ads and content (from fair pay start)", 3.0, "Rp m / month", "Your plan - confirm",
     "On top of the software and ads budget."),
    ("eap2_start", "New EAP contract 2 starts (psychologist starts the same month)", "Apr-27", "month", "Your plan - confirm",
     "Each EAP contract: Rp15m a month, cost Rp5m."),
    ("eap3_start", "New EAP contract 3 starts", "Jul-27", "month", "Your plan - confirm", "Must match a month label."),
    ("cof_cost", "Sales co-founder cash pay", 8.0, "Rp m / month", "Estimate", "Plus shares that vest over 4 years."),
    ("cof_start", "Sales co-founder starts", "Jul-27", "month", "Your plan - confirm", "Start only when a bank, OJK or LPS offer is open."),
    ("bank_start", "Bank agreement: first training day", "Oct-27", "month", "Your plan - confirm",
     "Bank days are priced and costed as bank events (Rate_Card)."),
    ("bank_days", "Bank training days per month from that month", 1, "days", "Your plan - confirm", "Set 0 to test 'no bank deal'."),
]
r0 = 100
first_input, sect_hdr = I["C98"], I["A84"]
ADDR = {}
for k, row in enumerate(rows):
    r = r0 + k
    if row[1] is None:
        I.cell(r, 1, row[0]); style_like(I.cell(r, 1), sect_hdr)
        continue
    key, item, val, unit, status, src = row
    for c, v in enumerate([key, item, val, unit, status, src], 1):
        cell = I.cell(r, c, v)
        style_like(cell, I.cell(86 if isinstance(val, str) else 97, c))
    if unit == "%":
        I.cell(r, 3).number_format = "0%"
    elif unit in ("switch", "days"):
        I.cell(r, 3).number_format = "0"
    elif unit != "month":
        I.cell(r, 3).number_format = "0.00"
    ADDR[key] = f"C{r}"
A = lambda k: "Inputs!$C$" + ADDR[k][1:]

# ------------------------------------------------------------------ 2. Volume: timing flag, new EAP, bank days
hdr = V["A22"]
V["A33"] = "Timing and growth plan (v14)"; style_like(V["A33"], hdr)
lab = [(34, "Pipeline gap: no new sales yet (1 = yes)", "flag", "MAX"),
       (35, "New EAP contracts active (growth plan)", "contracts", "MAX"),
       (36, "Bank training days (growth plan)", "days", "SUM")]
for r, t, u, agg in lab:
    V.cell(r, 1, t); V.cell(r, 2, u)
    for c in (1, 2):
        style_like(V.cell(r, c), V.cell(29, c))
    V.cell(r, 3, f"={agg}(E{r}:G{r})"); V.cell(r, 4, f"={agg}(H{r}:S{r})")
    style_like(V.cell(r, 3), V["C29"]); style_like(V.cell(r, 4), V["D29"])
for c in COLS:
    L = CL(c)
    V[f"{L}34"] = f"=IF(AND({A('pipe_on')}=1,{L}6<IFERROR(MATCH({A('sales_from')},$E$4:$S$4,0),1)),1,0)"
    V[f"{L}35"] = (f"={A('gp_on')}*(IF({L}6>=IFERROR(MATCH({A('eap2_start')},$E$4:$S$4,0),99),1,0)"
                   f"+IF({L}6>=IFERROR(MATCH({A('eap3_start')},$E$4:$S$4,0),99),1,0))")
    V[f"{L}36"] = f"={A('gp_on')}*IF({L}6>=IFERROR(MATCH({A('bank_start')},$E$4:$S$4,0),99),{A('bank_days')},0)*(1-{L}34)"
    for r in (34, 35, 36):
        style_like(V[f"{L}{r}"], V[f"{L}29"])
    V[f"{L}23"] = V[f"{L}23"].value + f"*(1-{L}34)"
    V[f"{L}24"] = V[f"{L}24"].value + f"*(1-{L}34)"
    V[f"{L}25"] = V[f"{L}25"].value + f"*(1-{L}34)+{L}36"
    V[f"{L}26"] = f"={L}18+{L}35"
    V[f"{L}27"] = f"={L}19*(1-{L}34*{A('pipe_b2c')})"
V["A25"] = "Bank KCP events (incl. growth-plan bank days)"
V["A26"] = "Bank EAP contracts active (existing + growth plan)"
V["A31"] = ("Notes: v14 adds the timing switch (no new sales before the first delivery month) and the growth plan "
            "(new EAP contracts and bank days). Change them in Inputs sections 11 and 12.")

# ------------------------------------------------------------------ 3. Capacity_OPEX: growth-plan fixed costs
C["A34"] = "Growth plan fixed costs (v14)"; style_like(C["A34"], C["A21"])
crow = [(35, "Fair pay active? (1 = yes)", "flag", "MAX"),
        (36, "Full-time psychologist active? (starts with EAP contract 2)", "flag", "MAX"),
        (37, "Fair-pay uplift: admin, social media, lead marketing", "Rp m", "SUM"),
        (38, "Fair-pay uplift: full-time psychologist", "Rp m", "SUM"),
        (39, "Extra IG ads and content", "Rp m", "SUM"),
        (40, "Sales co-founder", "Rp m", "SUM"),
        (41, "Growth plan fixed cost", "Rp m", "SUM")]
for r, t, u, agg in crow:
    C.cell(r, 1, t); C.cell(r, 2, u)
    src = 32 if r == 41 else (17 if u == "flag" else 26)
    for c in range(1, 20):
        style_like(C.cell(r, c), C.cell(src, c))
    C.cell(r, 3, f"={agg}(E{r}:G{r})"); C.cell(r, 4, f"={agg}(H{r}:S{r})")
for c in COLS:
    L = CL(c)
    C[f"{L}35"] = f"={A('gp_on')}*{A('fair_on')}*IF(Volume!{L}6>={month(ADDR['fair_start'])},1,0)"
    C[f"{L}36"] = f"={A('gp_on')}*{A('fair_on')}*IF(Volume!{L}6>={month(ADDR['eap2_start'])},1,0)"
    C[f"{L}37"] = (f"={L}35*(({A('fp_admin')}+{A('fp_socmed')}+{A('fp_lead')})*(1+{A('oncost')})"
                   f"-(Inputs!$C$43+Inputs!$C$40+Inputs!$C$39))")
    C[f"{L}38"] = f"={L}36*({A('fp_psych')}*(1+{A('oncost')})-Inputs!$C$42)"
    C[f"{L}39"] = f"={L}35*{A('ads_extra')}"
    C[f"{L}40"] = f"={A('gp_on')}*IF(Volume!{L}6>={month(ADDR['cof_start'])},{A('cof_cost')},0)"
    C[f"{L}41"] = f"=SUM({L}37:{L}40)"
    C[f"{L}32"] = f"=SUM({L}22:{L}31)+{L}41"

# ------------------------------------------------------------------ 4. PnL: route the new costs into the fixed lines
for c in COLS:
    L = CL(c)
    P[f"{L}27"] = P[f"{L}27"].value + f"+Capacity_OPEX!{L}37+Capacity_OPEX!{L}38+Capacity_OPEX!{L}39"
    P[f"{L}29"] = P[f"{L}29"].value + f"+Capacity_OPEX!{L}40"
P["A27"] = "Base: team, software, office, accountant (+ fair pay and ads)"
P["A29"] = "B2B sales (co-founder, Account Manager, travel)"

# ------------------------------------------------------------------ 5. Outlook sheet + combined chart
O = wb.create_sheet("Outlook", 1)
O["A1"] = "Outlook: sales, costs and cash each month (Rp million)"
O["A1"].font = Font(name="Arial", bold=True, size=14)
O["A2"] = ("Read the bars against the line: when the grey cost bar is taller than the blue sales bar, cash goes down. "
           "Cash is paid later than sales, because clients pay 0 to 60 days after the event.")
O["A2"].font = Font(name="Arial", italic=True, size=10)
lines = [("Month", None), ("Sales (revenue)", "=PnL!{L}12"), ("Delivery cost", "=PnL!{L}21"), ("Fixed cost", "=PnL!{L}31"),
         ("Tax", "=PnL!{L}34"), ("Total cost", "={O}6+{O}7+{O}8"), ("Net profit", "=PnL!{L}35"),
         ("Cash in", "=CashFlow!{L}11"), ("Cash out", "=CashFlow!{L}17"), ("Cash at end of month", "=CashFlow!{L}21"),
         ("Cash floor (1 month of costs)", "=CashFlow!{L}24"), ("Above the floor? (1 = yes)", "=CashFlow!{L}25")]
for k, (lab_, f) in enumerate(lines):
    r = 4 + k
    O.cell(r, 1, lab_).font = Font(name="Arial", bold=(k in (0, 1, 5, 9)), size=10)
    for j, c in enumerate(COLS):
        L = CL(c)
        cell = O.cell(r, 2 + j, f"=Volume!{L}4" if f is None else f.format(L=L, O=CL(2 + j)))
        cell.font = Font(name="Arial", size=10, bold=(k == 0), color="008000")
        cell.number_format = "0" if k == 11 else "#,##0"
O.column_dimensions["A"].width = 30
for j in range(15):
    O.column_dimensions[CL(2 + j)].width = 8
bar = BarChart(); bar.type = "col"; bar.grouping = "clustered"
bar.add_data(Reference(O, min_col=1, max_col=16, min_row=5, max_row=5), from_rows=True, titles_from_data=True)
bar.add_data(Reference(O, min_col=1, max_col=16, min_row=9, max_row=9), from_rows=True, titles_from_data=True)
bar.set_categories(Reference(O, min_col=2, max_col=16, min_row=4, max_row=4))
bar.series[0].graphicalProperties.solidFill = "576C94"
bar.series[1].graphicalProperties.solidFill = "B8BEC9"
ln = LineChart()
ln.add_data(Reference(O, min_col=1, max_col=16, min_row=13, max_row=13), from_rows=True, titles_from_data=True)
ln.add_data(Reference(O, min_col=1, max_col=16, min_row=14, max_row=14), from_rows=True, titles_from_data=True)
ln.series[0].graphicalProperties.line.solidFill = "475571"; ln.series[0].graphicalProperties.line.width = 38100
ln.series[1].graphicalProperties.line.solidFill = "9B3809"; ln.series[1].graphicalProperties.line.dashStyle = "dash"
ln.y_axis.axId = bar.y_axis.axId
bar.title = "Sales and total cost (bars) vs cash in the bank (line), Rp million"
bar.y_axis.title = "Rp million"
bar.y_axis.majorGridlines = None
for sr in ln.series:
    sr.smooth = False
bar.x_axis.delete = False
bar.y_axis.delete = False
bar.x_axis.title = None
bar += ln
bar.height, bar.width = 11, 30
bar.legend.position = "t"
O.add_chart(bar, "A18")
O.page_setup.orientation = "landscape"
O.page_setup.fitToWidth, O.page_setup.fitToHeight = 1, 1
O.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------------ 6. To_Fill: list the new estimates
last = T.max_row
new_fill = [("Timing", "First month new contracts are delivered", ADDR["sales_from"], "Decides how long we live on cash only."),
            ("Growth plan", "BPJS + THR on top of salary", ADDR["oncost"], "Fair-pay cost."),
            ("Growth plan", "Full-time psychologist gross pay", ADDR["fp_psych"], "Hiring step 2."),
            ("Growth plan", "EAP contract 2 and 3 start months", ADDR["eap2_start"], "Main source of new monthly income."),
            ("Growth plan", "Sales co-founder cash pay and start", ADDR["cof_cost"], "Hiring step 3."),
            ("Growth plan", "Bank training days per month and start", ADDR["bank_days"], "The swing factor for 2027.")]
for k, (pil, item, addr, why) in enumerate(new_fill):
    r = last + 1 + k
    vals = [int(T.cell(last, 1).value) + 1 + k, pil, item, f"=Inputs!{addr}", "", why, f"Inputs!{addr}", None, "Open"]
    for c, v in enumerate(vals, 1):
        T.cell(r, c, v); style_like(T.cell(r, c), T.cell(last, c))

# ------------------------------------------------------------------ 7. Summary: scenario snapshot from the Excel runs
RUNS_JSON = ROOT / "output" / "v14_runs.json"           # written by excel_runner.py; rebuild v14 after a new run
if RUNS_JSON.exists():
    import json
    runs = json.loads(RUNS_JSON.read_text())
    cols = [("Old timing, today's pay", "v13_timing_today_pay"), ("Real timing, today's pay", "real_timing_today_pay"),
            ("Plan, founders paid from Jan", "plan_founders_jan"), ("Plan, founders from May (ours)", "plan_founders_may"),
            ("Our plan, BI buys 30% less", "plan_founders_may_bi_cut")]
    S["A34"] = "Scenario comparison (v14 values from Excel; refresh with analysis/excel_runner.py, then rebuild)"
    mon = ["Oct-26", "Nov-26", "Dec-26"] + [m + "-27" for m in "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()]
    for j, (lab_, k) in enumerate(cols):
        f = runs[k]["fy27"]
        c = 2 + j
        S.cell(35, c, lab_)
        for r, v in [(36, f["R"]), (37, f["N"]), (38, f["m"]), (39, f["low"]), (40, mon[f["low_idx"]]),
                     (41, f["below_floor"]), (42, f["end"])]:
            S.cell(r, c, v)
    S["A43"], S["A44"] = None, None
    for c in range(2, 7):
        S.cell(43, c).value = None

wb.save(OUT)
print("wrote", OUT.name, "inputs:", ", ".join(f"{k}={v}" for k, v in ADDR.items()))
