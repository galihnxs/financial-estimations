"""Builds the internal RAB workbook for the BI Fit and Proper Test program.

Run:  python build_rab.py  (needs openpyxl). Output lands one folder up.
Blue cells on "Asumsi" are inputs. Every other number is a formula.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "RAB - BI Fit and Proper Test (Internal).xlsx"

RP = '"Rp"#,##0;[Red]-"Rp"#,##0'
PCT = "0.0%"
BLUE = Font(name="Calibri", color="0000FF")
BOLD = Font(name="Calibri", bold=True)
WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="475571")
SUB_FILL = PatternFill("solid", fgColor="E4E8F0")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="B7BFCF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = WHITE_BOLD, HEAD_FILL, WRAP, BOX


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------------------------------------------------------- Asumsi
asm = wb.active
asm.title = "Asumsi"
asm["A1"] = "ASUMSI - BI Fit and Proper Test (INTERNAL, bukan untuk klien). Sel biru = input."
asm["A1"].font = Font(name="Calibri", bold=True, size=13)
header(asm, 2, ["Asumsi", "Nilai", "Satuan", "Sumber / catatan"])

# name -> (label, value, unit, source). Order sets the row.
ASSUMPTIONS = [
    ("peserta", "Jumlah peserta", 6, "orang", "Penawaran 28 Okt 2026, disetujui BI"),
    ("harga", "Harga final disepakati BI (total)", 15_000_000, "Rp", 'BI: "pakai Rp15.000.000 saja". Naik dari penawaran Rp14.520.000'),
    ("pkp", "Status pajak Teman Journey (PKP / Non-PKP)", "Non-PKP", "-", "Dikonfirmasi Galih 28 Okt 2026: Non-PKP. Ganti ke PKP hanya jika sudah dikukuhkan PKP"),
    ("ppn", "Tarif PPN efektif (hanya jika PKP)", 0.11, "%", "12% x DPP nilai lain 11/12 = 11% efektif"),
    ("pph23", "PPh 23 dipotong BI", 0.02, "%", "Sama dengan model BI Sibolga. Bisa dikreditkan di SPT badan, di sini dihitung sebagai biaya (konservatif)"),
    ("tarif_dok", "Honor psikolog: penyusunan dokumen per jam", 110_000, "Rp/jam", "Diberikan Galih (28 Okt 2026): Rp110.000 per jam untuk pembuatan dokumen"),
    ("jam_rubrik", "Jam minimal per psikolog per rubrik", 2, "jam", "Diberikan Galih: tiap psikolog minimal 2 jam"),
    ("psi_rubrik", "Psikolog yang menyusun tiap rubrik", 2, "orang", "Juri utama + psikolog fasilitator LGD menyusun bersama"),
    ("tarif_tatap", "Honor psikolog: tatap muka asesmen per jam (biaya dasar konseling)", 350_000, "Rp/jam", "Biaya dasar konseling TJ: honor psikolog Rp350.000 per sesi 60 menit (BI Sibolga, dari Kia)"),
    ("jam_koordinasi", "Jam koordinasi kebutuhan BI per psikolog", 1, "jam", "Kick-off: kompetensi target, jadwal, format laporan"),
    ("jam_kasus", "Jam penyusunan studi kasus dan instruksi LGD", 2, "jam", "Disusun psikolog fasilitator LGD"),
    ("jam_lgd", "Durasi LGD termasuk briefing", 2, "jam", "Briefing 15 menit, diskusi 60-75 menit, penutupan"),
    ("psi_lgd", "Psikolog di sesi LGD", 2, "orang", "1 fasilitator + 1 asesor, masing-masing mengamati 3 peserta"),
    ("jam_skor_lgd", "Jam skoring LGD per psikolog", 1, "jam", "Mengisi rubrik dan catatan perilaku"),
    ("jam_wawancara", "Durasi wawancara per peserta", 1, "jam", "Wawancara terstruktur berbasis perilaku"),
    ("jam_skor_wwc", "Jam skoring dan catatan wawancara per peserta", 0.5, "jam", ""),
    ("jam_laporan", "Jam penulisan laporan per peserta", 2, "jam", "Integrasi tes, LGD, wawancara"),
    ("jam_review", "Jam review mutu laporan oleh juri utama per peserta", 0.5, "jam", ""),
    ("jam_integrasi", "Jam rapat integrasi hasil per psikolog", 1, "jam", "Menyepakati rekomendasi akhir tiap peserta"),
    ("tes", "Tes BFI + DISC per peserta", 120_000, "Rp/peserta", "RAB lama"),
    ("admin", "Admin dan koordinasi jadwal per peserta", 35_000, "Rp/peserta", "Rujukan admin per sesi BI Sibolga"),
    ("transport", "Transport psikolog per orang per hari", 150_000, "Rp", "ASUMSI: pelaksanaan tatap muka di Jakarta. Isi 0 jika online"),
    ("hari", "Total kehadiran psikolog (orang x hari)", 3, "orang-hari", "ASUMSI: juri utama 2 hari (LGD + wawancara), fasilitator 1 hari"),
    ("kontrak", "Kontrak dan materai", 150_000, "Rp", "Rujukan BI Sibolga"),
    ("komisi", "Komisi (% dari profit kotor)", 0.10, "%", "RAB lama"),
    ("opex", "Biaya operasional bulanan TJ", 20_000_000, "Rp", "Rujukan BI Sibolga"),
    ("target", "Target profit nett", 1_000_000, "Rp", "RAB lama"),
]
A = {}  # name -> absolute reference on Asumsi
for r, (name, label, value, unit, src) in enumerate(ASSUMPTIONS, start=3):
    asm.cell(row=r, column=1, value=label)
    c = asm.cell(row=r, column=2, value=value)
    c.font = BLUE
    if unit == "%":
        c.number_format = PCT
    elif unit.startswith("Rp"):
        c.number_format = RP
    asm.cell(row=r, column=3, value=unit)
    asm.cell(row=r, column=4, value=src).alignment = WRAP
    for col in range(1, 5):
        asm.cell(row=r, column=col).border = BOX
    A[name] = f"Asumsi!$B${r}"
widths(asm, [48, 16, 12, 80])

# ---------------------------------------------------------------- RAB
rab = wb.create_sheet("RAB")
rab["A1"] = "RAB INTERNAL - Asesmen Fit and Proper Test, Bank Indonesia (6 peserta)"
rab["A1"].font = Font(name="Calibri", bold=True, size=13)
rab["A2"] = "Biaya psikolog = orang x jam x kali x tarif per jam. Kolom Komponen menghubungkan tiap biaya ke layanan di proposal."
header(rab, 4, ["No", "Kategori", "Item biaya", "Orang", "Jam / orang", "Kali", "Qty", "Satuan", "Tarif", "Total", "Komponen proposal", "Jenis"])

P = A["peserta"]
# (kategori, item, orang, jam, kali, satuan, tarif, komponen, jenis)
LINES = [
    ("Persiapan", "Koordinasi kebutuhan dan kompetensi target dengan BI", A["psi_rubrik"], A["jam_koordinasi"], 1, "jam", A["tarif_dok"], 1, "Honor"),
    ("Persiapan", "Rubrik penilaian wawancara", A["psi_rubrik"], A["jam_rubrik"], 1, "jam", A["tarif_dok"], 1, "Honor"),
    ("Persiapan", "Rubrik penilaian LGD", A["psi_rubrik"], A["jam_rubrik"], 1, "jam", A["tarif_dok"], 1, "Honor"),
    ("Persiapan", "Studi kasus dan instruksi LGD", 1, A["jam_kasus"], 1, "jam", A["tarif_dok"], 1, "Honor"),
    ("Tes", "Tes BFI + DISC (platform dan skoring)", 1, 1, P, "peserta", A["tes"], 2, "Langsung"),
    ("LGD", "Pelaksanaan LGD (fasilitator + asesor)", A["psi_lgd"], A["jam_lgd"], 1, "jam", A["tarif_tatap"], 3, "Honor"),
    ("LGD", "Skoring dan catatan perilaku LGD", A["psi_lgd"], A["jam_skor_lgd"], 1, "jam", A["tarif_dok"], 3, "Honor"),
    ("Wawancara", "Wawancara oleh psikolog juri utama", 1, A["jam_wawancara"], P, "jam", A["tarif_tatap"], 4, "Honor"),
    ("Wawancara", "Skoring dan catatan wawancara", 1, A["jam_skor_wwc"], P, "jam", A["tarif_dok"], 4, "Honor"),
    ("Laporan", "Penulisan laporan asesmen komprehensif", 1, A["jam_laporan"], P, "jam", A["tarif_dok"], 5, "Honor"),
    ("Laporan", "Review mutu laporan oleh juri utama", 1, A["jam_review"], P, "jam", A["tarif_dok"], 5, "Honor"),
    ("Laporan", "Rapat integrasi hasil dan rekomendasi", A["psi_rubrik"], A["jam_integrasi"], 1, "jam", A["tarif_dok"], 5, "Honor"),
    ("Operasional", "Admin dan koordinasi jadwal peserta", 1, 1, P, "peserta", A["admin"], "Overhead", "Langsung"),
    ("Operasional", "Transport psikolog", 1, 1, A["hari"], "orang-hari", A["transport"], "Overhead", "Langsung"),
    ("Operasional", "Kontrak dan materai", 1, 1, 1, "deal", A["kontrak"], "Overhead", "Langsung"),
]
first = 5
for i, (kat, item, orang, jam, kali, sat, tarif, komp, jenis) in enumerate(LINES):
    r = first + i
    vals = [i + 1, kat, item, f"={orang}" if isinstance(orang, str) else orang,
            f"={jam}" if isinstance(jam, str) else jam, f"={kali}" if isinstance(kali, str) else kali,
            f"=D{r}*E{r}*F{r}", sat, f"={tarif}", f"=G{r}*I{r}", komp, jenis]
    for col, v in enumerate(vals, start=1):
        c = rab.cell(row=r, column=col, value=v)
        c.border = BOX
    rab.cell(row=r, column=9).number_format = RP
    rab.cell(row=r, column=10).number_format = RP
last = first + len(LINES) - 1
TOT = last + 1
rab.cell(row=TOT, column=9, value="Jumlah Biaya").font = BOLD
rab.cell(row=TOT, column=10, value=f"=SUM(J{first}:J{last})").font = BOLD
rab.cell(row=TOT, column=10).number_format = RP
rab.cell(row=TOT + 1, column=9, value="Biaya per peserta")
rab.cell(row=TOT + 1, column=10, value=f"=J{TOT}/{P}").number_format = RP
rab.cell(row=TOT + 2, column=9, value="Total honor psikolog")
rab.cell(row=TOT + 2, column=10, value=f'=SUMIF(L{first}:L{last},"Honor",J{first}:J{last})').number_format = RP
COST = f"RAB!$J${TOT}"
HONOR = f"RAB!$J${TOT + 2}"

# Revenue and profit block
r0 = TOT + 4
rab.cell(row=r0, column=2, value="PENDAPATAN DAN PROFIT").font = BOLD
rows = [
    ("NILAI AKHIR (harga disepakati BI)", f"={A['harga']}", RP),
    ("PPN terutang (0 jika Non-PKP)", f'=IF({A["pkp"]}="PKP",B{{r1}}-B{{r1}}/(1+{A["ppn"]}),0)', RP),
    ("Pendapatan bersih PPN (DPP)", "=B{r1}-B{r2}", RP),
    ("PPh 23 dipotong BI", f"=B{{r3}}*{A['pph23']}", RP),
    ("AFTER PAJAK", "=B{r3}-B{r4}", RP),
    ("Jumlah Biaya", f"={COST}", RP),
    ("AFTER COST (profit kotor)", "=B{r5}-B{r6}", RP),
    ("Komisi", f"=MAX(0,B{{r7}}*{A['komisi']})", RP),
    ("PROFIT NETT", "=B{r7}-B{r8}", RP),
    ("Margin nett terhadap harga", "=B{r9}/B{r1}", PCT),
    ("Markup harga (DPP) terhadap biaya", f"=B{{r3}}/{COST}-1", PCT),
    ("Capai target profit nett?", f'=IF(B{{r9}}>={A["target"]},"Ya","Belum")', None),
    ("AFTER COST jika semua sesi oleh psikolog in-house (honor Rp0)", f"=B{{r7}}+{HONOR}", RP),
    ("% biaya operasional 1 bulan tertutup (profit nett)", f"=B{{r9}}/{A['opex']}", PCT),
]
refs = {f"r{k}": r0 + k for k in range(1, len(rows) + 1)}
for k, (label, formula, fmt) in enumerate(rows, start=1):
    r = r0 + k
    rab.cell(row=r, column=1, value=label)
    c = rab.cell(row=r, column=2, value=formula.format(**refs))
    if fmt:
        c.number_format = fmt
    if label.isupper() or label.startswith(("NILAI", "AFTER", "PROFIT")):
        rab.cell(row=r, column=1).font = BOLD
        c.font = BOLD
    if label.startswith("PROFIT NETT"):
        rab.cell(row=r, column=1).fill = KEY_FILL
        c.fill = KEY_FILL
PROFIT = f"RAB!$B${refs['r9']}"
MARGIN = f"RAB!$B${refs['r10']}"
MARKUP = f"RAB!$B${refs['r11']}"
AFTER_TAX = f"RAB!$B${refs['r5']}"
INHOUSE = f"RAB!$B${refs['r13']}"

# Markup scenarios, same shape as the old RAB and the Sibolga model
m0 = r0 + len(rows) + 2
rab.cell(row=m0, column=1, value="SKENARIO MARKUP (cek: harga Rp15jt ada di markup berapa)").font = BOLD
header_cells = ["Markup", "Harga jual total", "Harga / peserta", "Profit nett (setelah pajak, komisi)"]
for i, label in enumerate(header_cells, start=1):
    c = rab.cell(row=m0 + 1, column=i, value=label)
    c.font, c.fill, c.border, c.alignment = WHITE_BOLD, HEAD_FILL, BOX, WRAP
for j, mk in enumerate([0.3, 0.4, 0.5, 0.75, 1.0]):
    r = m0 + 2 + j
    rab.cell(row=r, column=1, value=mk).number_format = PCT
    rab.cell(row=r, column=2, value=f"={COST}*(1+A{r})").number_format = RP
    rab.cell(row=r, column=3, value=f"=B{r}/{P}").number_format = RP
    # Non-PKP view: price minus PPh 23, minus cost, minus commission
    rab.cell(row=r, column=4, value=f"=(B{r}*(1-{A['pph23']})-{COST})*(1-{A['komisi']})").number_format = RP
widths(rab, [44, 18, 46, 8, 11, 8, 8, 11, 14, 15, 18, 10])
for r in range(first, last + 1):
    rab.cell(row=r, column=3).alignment = WRAP

# ---------------------------------------------------------------- Harga Proposal
hp = wb.create_sheet("Harga Proposal")
hp["A1"] = "HARGA PER LAYANAN UNTUK PROPOSAL (harga jual, bukan biaya). Sel biru = harga yang ditulis di proposal."
hp["A1"].font = Font(name="Calibri", bold=True, size=13)
hp["A2"] = "Total harus sama dengan harga disepakati BI. Overhead dibagi ke tiap layanan sesuai porsi biaya langsungnya."
header(hp, 4, ["No", "Layanan (bahasa proposal)", "Volume", "Satuan", "Harga satuan", "Jumlah", "Biaya langsung", "Biaya + porsi overhead", "Markup"])
SERVICES = [
    (1, "Penyesuaian kebutuhan dan penyusunan rubrik penilaian", 1, "paket", 2_800_000),
    (2, "Tes psikologi BFI dan DISC", f"={P}", "peserta", 200_000),
    (3, "Leaderless Group Discussion (LGD)", 1, "sesi", 2_750_000),
    (4, "Wawancara oleh psikolog juri utama", f"={P}", "peserta", 675_000),
    (5, "Laporan asesmen komprehensif dan rekomendasi", f"={P}", "peserta", 700_000),
]
s0 = 5
rng_k = f"RAB!$K${first}:$K${last}"
rng_j = f"RAB!$J${first}:$J${last}"
for i, (no, name, vol, sat, price) in enumerate(SERVICES):
    r = s0 + i
    hp.cell(row=r, column=1, value=no)
    hp.cell(row=r, column=2, value=name).alignment = WRAP
    hp.cell(row=r, column=3, value=vol)
    hp.cell(row=r, column=4, value=sat)
    c = hp.cell(row=r, column=5, value=price)
    c.font, c.number_format = BLUE, RP
    hp.cell(row=r, column=6, value=f"=C{r}*E{r}").number_format = RP
    hp.cell(row=r, column=7, value=f"=SUMIF({rng_k},A{r},{rng_j})").number_format = RP
    hp.cell(row=r, column=8, value=f'=G{r}+G{r}/SUM($G${s0}:$G${s0 + 4})*SUMIF({rng_k},"Overhead",{rng_j})').number_format = RP
    hp.cell(row=r, column=9, value=f"=F{r}/H{r}-1").number_format = PCT
    for col in range(1, 10):
        hp.cell(row=r, column=col).border = BOX
ht = s0 + len(SERVICES)
hp.cell(row=ht, column=2, value="TOTAL").font = BOLD
for col in (6, 7, 8):
    L = get_column_letter(col)
    c = hp.cell(row=ht, column=col, value=f"=SUM({L}{s0}:{L}{ht - 1})")
    c.font, c.number_format = BOLD, RP
hp.cell(row=ht, column=9, value=f"=F{ht}/H{ht}-1").number_format = PCT
hp.cell(row=ht + 1, column=2, value="Harga per peserta")
hp.cell(row=ht + 1, column=6, value=f"=F{ht}/{P}").number_format = RP
hp.cell(row=ht + 2, column=2, value="Cek: total = harga disepakati BI")
hp.cell(row=ht + 2, column=6, value=f'=IF(F{ht}={A["harga"]},"OK","TIDAK SAMA")').font = BOLD
widths(hp, [5, 52, 9, 10, 15, 15, 15, 20, 10])

# ---------------------------------------------------------------- RAB Lama vs Baru
cmp_ = wb.create_sheet("RAB Lama vs Baru")
cmp_["A1"] = "Kenapa RAB lama tidak bisa dipakai: total biayanya Rp15.020.000, lebih besar dari harga Rp15.000.000"
cmp_["A1"].font = Font(name="Calibri", bold=True, size=13)
header(cmp_, 3, ["Item RAB lama", "RAB lama", "Baris RAB baru", "RAB baru", "Selisih", "Kenapa berubah"])


def rab_sum(nos):
    return "=" + "+".join(f"RAB!$J${first + n - 1}" for n in nos)


COMPARE = [
    ("Psikolog juri utama (Rp750.000 x 6 kandidat)", 4_500_000, "8-9", rab_sum([8, 9]),
     "Lama: tarif borongan per kandidat tanpa dasar jam. Baru: 1 jam wawancara (Rp350rb/jam) + 0,5 jam skoring (Rp110rb/jam) per peserta"),
    ("Psikolog LGD (Rp2.000.000 per sesi)", 2_000_000, "6-7", rab_sum([6, 7]),
     "Baru: 2 psikolog x 2 jam tatap muka + 1 jam skoring masing-masing"),
    ("BFI - DISC (Rp120.000 x 6)", 720_000, "5", rab_sum([5]), "Tidak berubah"),
    ("Rubrik asesmen LGD (Rp2.500.000)", 2_500_000, "3-4", rab_sum([3, 4]),
     "Baru: 2 psikolog x 2 jam x Rp110rb/jam, ditambah studi kasus LGD 2 jam"),
    ("Rubrik asesmen wawancara (Rp2.750.000)", 2_750_000, "1-2", rab_sum([1, 2]),
     "Baru: 2 psikolog x 2 jam x Rp110rb/jam, ditambah 1 jam koordinasi kebutuhan BI per psikolog"),
    ("Laporan asesmen komprehensif (Rp425.000 x 6)", 2_550_000, "10-12", rab_sum([10, 11, 12]),
     "Baru: 2 jam tulis + 0,5 jam review per laporan, ditambah 1 jam rapat integrasi per psikolog"),
    ("(tidak ada)", 0, "13-15", rab_sum([13, 14, 15]),
     "Baru: admin, transport, kontrak dan materai yang dulu belum dihitung"),
]
for i, (old_item, old_val, rows_new, new_f, why) in enumerate(COMPARE):
    r = 4 + i
    vals = [old_item, old_val, rows_new, new_f, f"=D{r}-B{r}", why]
    for col, v in enumerate(vals, start=1):
        c = cmp_.cell(row=r, column=col, value=v)
        c.border, c.alignment = BOX, WRAP
    for col in (2, 4, 5):
        cmp_.cell(row=r, column=col).number_format = RP
ct = 4 + len(COMPARE)
cmp_.cell(row=ct, column=1, value="TOTAL BIAYA").font = BOLD
for col in (2, 4, 5):
    L = get_column_letter(col)
    c = cmp_.cell(row=ct, column=col, value=f"=SUM({L}4:{L}{ct - 1})")
    c.font, c.number_format = BOLD, RP
cmp_.cell(row=ct + 1, column=1, value="Profit kotor pada harga Rp15jt (sebelum pajak dan komisi)")
for col in (2, 4):
    L = get_column_letter(col)
    cmp_.cell(row=ct + 1, column=col, value=f"={A['harga']}-{L}{ct}").number_format = RP
widths(cmp_, [44, 15, 12, 15, 15, 80])

# ---------------------------------------------------------------- Ringkasan (first tab)
rk = wb.create_sheet("Ringkasan", 0)
rk["A1"] = "RINGKASAN - BI Fit and Proper Test (INTERNAL, bukan untuk klien)"
rk["A1"].font = Font(name="Calibri", bold=True, size=13)
SUMMARY = [
    ("Jumlah peserta", f"={P}", "0"),
    ("Harga disepakati BI", f"={A['harga']}", RP),
    ("Harga per peserta", f"={A['harga']}/{P}", RP),
    ("Status pajak", f"={A['pkp']}", None),
    ("Diterima setelah pajak (PPN jika PKP, PPh 23)", f"={AFTER_TAX}", RP),
    ("Jumlah biaya", f"={COST}", RP),
    ("Biaya per peserta", f"={COST}/{P}", RP),
    ("Markup harga terhadap biaya", f"={MARKUP}", PCT),
    ("PROFIT NETT (setelah pajak dan komisi)", f"={PROFIT}", RP),
    ("Margin nett terhadap harga", f"={MARGIN}", PCT),
    ("AFTER COST jika semua sesi oleh psikolog in-house", f"={INHOUSE}", RP),
    ("Cek harga per layanan di proposal", f"='Harga Proposal'!F{ht + 2}", None),
]
for i, (label, f, fmt) in enumerate(SUMMARY, start=3):
    rk.cell(row=i, column=1, value=label).border = BOX
    c = rk.cell(row=i, column=2, value=f)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if label.startswith("PROFIT"):
        for col in (1, 2):
            rk.cell(row=i, column=col).font = BOLD
            rk.cell(row=i, column=col).fill = KEY_FILL
notes_row = 3 + len(SUMMARY) + 1
NOTES = [
    "Catatan:",
    "1. Semua honor psikolog dihitung per jam: Rp110.000/jam untuk kerja dokumen (rubrik, skoring, laporan), Rp350.000/jam untuk tatap muka (LGD, wawancara) sesuai biaya dasar konseling.",
    "2. Rubrik disusun 2 psikolog, masing-masing minimal 2 jam, sesuai arahan Galih.",
    "3. TJ Non-PKP (dikonfirmasi 28 Okt 2026), jadi tidak ada PPN. Penawaran 28 Okt yang menulis 'sudah termasuk PPN' perlu dikoreksi jika dikirim ulang.",
    "4. Transport mengasumsikan tatap muka di Jakarta. Jika lokasi lain atau online, ubah di Asumsi.",
    "5. Tab 'Harga Proposal' adalah satu-satunya angka yang boleh keluar ke BI. Tab lain internal.",
]
for i, n in enumerate(NOTES):
    c = rk.cell(row=notes_row + i, column=1, value=n)
    c.alignment = Alignment(wrap_text=False)
    if i == 0:
        c.font = BOLD
widths(rk, [52, 20])

wb.save(OUT)
print("saved", OUT)
