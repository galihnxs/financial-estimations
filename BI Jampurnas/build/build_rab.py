"""Builds the internal RAB workbook for BI JAMPURNAS 2026 (Jambore Pejuang Rupiah Nasional).

Run:  python build_rab.py  (needs openpyxl). Output lands one folder up.
Blue cells on "Asumsi" are inputs. Every other number is a formula.
Same layout as the BI Fit and Proper RAB. Costs are the fixed fees agreed on 28 Oct 2026.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "RAB - BI Jampurnas 2026 (Internal).xlsx"

RP = '"Rp"#,##0;[Red]-"Rp"#,##0'
PCT = "0.0%"
TIMES = '0.00"x"'
BLUE = Font(name="Calibri", color="0000FF")
BOLD = Font(name="Calibri", bold=True)
WHITE_BOLD = Font(name="Calibri", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="475571")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="B7BFCF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = WHITE_BOLD, HEAD_FILL, WRAP, BOX


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------------------------------------------------------- Asumsi
asm = wb.active
asm.title = "Asumsi"
asm["A1"] = "ASUMSI - BI JAMPURNAS 2026 (INTERNAL, bukan untuk klien). Sel biru = input."
asm["A1"].font = Font(name="Calibri", bold=True, size=13)
header(asm, 2, ["Asumsi", "Nilai", "Satuan", "Sumber / catatan"])

# name -> (label, value, unit, source). Order sets the row.
ASSUMPTIONS = [
    ("harga", "Harga yang diajukan ke BI (total)", 15_000_000, "Rp", "Galih 28 Okt 2026: ajukan Rp15.000.000"),
    ("pkp", "Status pajak Teman Journey (PKP / Non-PKP)", "Non-PKP", "-", "Dikonfirmasi 28 Okt 2026"),
    ("ppn", "Tarif PPN efektif (hanya jika PKP)", 0.11, "%", "12% x DPP nilai lain 11/12"),
    ("pph23", "PPh 23 dipotong BI", 0.02, "%", "Sama dengan RAB Fit and Proper. Dihitung sebagai biaya (konservatif)"),
    ("karya", "Jumlah karya yang dinilai", 46, "karya", "Galih 28 Okt 2026: fee juri sudah termasuk penilaian 46 karya"),
    ("fee_dosen", "Fee juri dosen public speaking (tetap)", 2_500_000, "Rp/juri",
     "Galih 28 Okt 2026: fee tetap, termasuk hari acara dan rubrik penilaian 46 karya. Tidak ada fee offline terpisah"),
    ("fee_psikolog", "Fee juri psikolog (tetap)", 2_500_000, "Rp/juri",
     "Galih 28 Okt 2026: fee tetap, termasuk hari acara dan rubrik penilaian 46 karya. Tidak ada fee offline terpisah"),
    ("fee_pm", "Fee project manager", 2_000_000, "Rp/proyek", "Galih 28 Okt 2026: Rp2 juta untuk proyek ini, sesuai kesepakatan"),
    ("admin", "Biaya administrasi (dibayar ke Coreitera)", 400_000, "Rp/proyek", "Galih 28 Okt 2026: Rp400rb ke Coreitera"),
    ("komisi", "Komisi (% dari profit kotor)", 0.10, "%", "Sama dengan RAB Fit and Proper"),
    ("margin_min", "Batas bawah margin nett", 0.35, "%", "Margin Fit and Proper 35,6% diterima"),
    ("margin_target", "Target margin nett", 0.40, "%", "Galih menanyakan cara mencapai 40% (28 Okt 2026)"),
]
A = {}
for r, (name, label, value, unit, src) in enumerate(ASSUMPTIONS, start=3):
    asm.cell(row=r, column=1, value=label)
    c = asm.cell(row=r, column=2, value=value)
    c.font = BLUE
    if unit == "%":
        c.number_format = PCT
    elif unit.startswith("Rp"):
        c.number_format = RP
    asm.cell(row=r, column=3, value=unit)
    asm.cell(row=r, column=4, value=src).alignment = WRAP
    for col in range(1, 5):
        asm.cell(row=r, column=col).border = BOX
    A[name] = f"Asumsi!$B${r}"
widths(asm, [52, 16, 12, 84])

# ---------------------------------------------------------------- RAB
rab = wb.create_sheet("RAB")
rab["A1"] = "RAB INTERNAL - BI JAMPURNAS 2026 (Kantor Pusat Bank Indonesia)"
rab["A1"].font = Font(name="Calibri", bold=True, size=13)
rab["A2"] = ("Tim 3 orang: juri dosen public speaking, juri psikolog, project manager. "
             "Semua fee tetap per proyek, sesuai kesepakatan 28 Okt 2026.")
header(rab, 4, ["No", "Kategori", "Item biaya", "Qty", "Satuan", "Tarif", "Total", "Komponen proposal", "Jenis"])

# (kategori, item, qty, satuan, tarif, komponen, jenis)
LINES = [
    ("Juri", "Juri dosen public speaking: hari acara + rubrik penilaian 46 karya", 1, "juri", A["fee_dosen"], 1, "Honor"),
    ("Juri", "Juri psikolog: hari acara + rubrik penilaian 46 karya", 1, "juri", A["fee_psikolog"], 2, "Honor"),
    ("Koordinasi", "Project manager", 1, "proyek", A["fee_pm"], 3, "Honor"),
    ("Operasional", "Biaya administrasi (Coreitera)", 1, "proyek", A["admin"], "Overhead", "Langsung"),
]
first = 5
for i, (kat, item, qty, sat, tarif, komp, jenis) in enumerate(LINES):
    r = first + i
    vals = [i + 1, kat, item, qty, sat, f"={tarif}", f"=D{r}*F{r}", komp, jenis]
    for col, v in enumerate(vals, start=1):
        rab.cell(row=r, column=col, value=v).border = BOX
    rab.cell(row=r, column=6).number_format = RP
    rab.cell(row=r, column=7).number_format = RP
    rab.cell(row=r, column=3).alignment = WRAP
last = first + len(LINES) - 1
TOT = last + 1
rab.cell(row=TOT, column=6, value="Jumlah Biaya").font = BOLD
rab.cell(row=TOT, column=7, value=f"=SUM(G{first}:G{last})").font = BOLD
rab.cell(row=TOT, column=7).number_format = RP
rab.cell(row=TOT + 1, column=6, value="Biaya juri per karya")
rab.cell(row=TOT + 1, column=7, value=f"=(G{first}+G{first + 1})/{A['karya']}").number_format = RP
COST = f"RAB!$G${TOT}"

r0 = TOT + 3
rab.cell(row=r0, column=1, value="PENDAPATAN DAN PROFIT").font = BOLD
rows = [
    ("NILAI AKHIR (harga diajukan ke BI)", f"={A['harga']}", RP),
    ("PPN terutang (0 jika Non-PKP)", f'=IF({A["pkp"]}="PKP",B{{r1}}-B{{r1}}/(1+{A["ppn"]}),0)', RP),
    ("Pendapatan bersih PPN (DPP)", "=B{r1}-B{r2}", RP),
    ("PPh 23 dipotong BI", f"=B{{r3}}*{A['pph23']}", RP),
    ("AFTER PAJAK", "=B{r3}-B{r4}", RP),
    ("Jumlah Biaya", f"={COST}", RP),
    ("AFTER COST (profit kotor)", "=B{r5}-B{r6}", RP),
    ("Komisi", f"=MAX(0,B{{r7}}*{A['komisi']})", RP),
    ("PROFIT NETT", "=B{r7}-B{r8}", RP),
    ("Margin nett terhadap harga", "=B{r9}/B{r1}", PCT),
    ("Markup harga (DPP) terhadap biaya", f"=B{{r3}}/{COST}-1", PCT),
    ("Harga / biaya", f"=B{{r1}}/{COST}", TIMES),
    ("Status margin", f'=IF(B{{r10}}>={A["margin_target"]},"Capai target 40%",IF(B{{r10}}>={A["margin_min"]},"Dalam rentang 35-40%","Di bawah 35%"))', None),
]
refs = {f"r{k}": r0 + k for k in range(1, len(rows) + 1)}
for k, (label, formula, fmt) in enumerate(rows, start=1):
    r = r0 + k
    rab.cell(row=r, column=1, value=label)
    c = rab.cell(row=r, column=2, value=formula.format(**refs))
    if fmt:
        c.number_format = fmt
    if label.startswith(("NILAI", "AFTER", "PROFIT")):
        rab.cell(row=r, column=1).font = BOLD
        c.font = BOLD
    if label.startswith(("PROFIT NETT", "Margin nett", "Status")):
        rab.cell(row=r, column=1).fill = KEY_FILL
        c.fill = KEY_FILL
PROFIT = f"RAB!$B${refs['r9']}"
MARGIN = f"RAB!$B${refs['r10']}"
AFTER_TAX = f"RAB!$B${refs['r5']}"

# Price needed for a target net margin: P = (1-komisi) x cost / ((1-pph23) x (1-komisi) - margin)
q0 = r0 + len(rows) + 2
rab.cell(row=q0, column=1, value="HARGA UNTUK MARGIN NETT TARGET (Non-PKP, setelah PPh 23 dan komisi)").font = BOLD
header_cells = ["Margin nett", "Harga ke BI", "Profit nett", "Harga / biaya"]
for i, label in enumerate(header_cells, start=1):
    c = rab.cell(row=q0 + 1, column=i, value=label)
    c.font, c.fill, c.border, c.alignment = WHITE_BOLD, HEAD_FILL, BOX, WRAP
PRICE_ROWS = {}
for i, m in enumerate([0.30, 0.35, 0.40]):
    r = q0 + 2 + i
    rab.cell(row=r, column=1, value=m).number_format = PCT
    rab.cell(row=r, column=2, value=f"=(1-{A['komisi']})*{COST}/((1-{A['pph23']})*(1-{A['komisi']})-A{r})").number_format = RP
    rab.cell(row=r, column=3, value=f"=A{r}*B{r}").number_format = RP
    rab.cell(row=r, column=4, value=f"=B{r}/{COST}").number_format = TIMES
    for col in range(1, 5):
        rab.cell(row=r, column=col).border = BOX
    PRICE_ROWS[m] = r
widths(rab, [44, 18, 58, 8, 10, 14, 15, 18, 10])

# ---------------------------------------------------------------- Harga Proposal (draft)
hp = wb.create_sheet("Harga Proposal")
hp["A1"] = "HARGA PER LAYANAN UNTUK PROPOSAL (DRAF, harga jual bukan biaya). Sel biru = harga yang ditulis di proposal."
hp["A1"].font = Font(name="Calibri", bold=True, size=13)
hp["A2"] = "Total harus sama dengan harga yang diajukan. Overhead dibagi ke tiap layanan sesuai porsi biaya langsungnya."
header(hp, 4, ["No", "Layanan (bahasa proposal)", "Volume", "Satuan", "Harga satuan", "Jumlah", "Biaya langsung", "Biaya + porsi overhead", "Harga / biaya"])
SERVICES = [
    (1, "Juri dosen public speaking: penilaian 46 karya dan penjurian di hari acara", 1, "paket", 6_000_000),
    (2, "Juri psikolog: penilaian 46 karya dan penjurian di hari acara", 1, "paket", 6_000_000),
    (3, "Manajemen program dan pendampingan acara", 1, "paket", 3_000_000),
]
p0 = 5
pn = p0 + len(SERVICES) - 1
rng_k = f"RAB!$H${first}:$H${last}"
rng_g = f"RAB!$G${first}:$G${last}"
for i, (no, name, vol, sat, price) in enumerate(SERVICES):
    r = p0 + i
    hp.cell(row=r, column=1, value=no)
    hp.cell(row=r, column=2, value=name).alignment = WRAP
    hp.cell(row=r, column=3, value=vol)
    hp.cell(row=r, column=4, value=sat)
    c = hp.cell(row=r, column=5, value=price)
    c.font, c.number_format = BLUE, RP
    hp.cell(row=r, column=6, value=f"=C{r}*E{r}").number_format = RP
    hp.cell(row=r, column=7, value=f"=SUMIF({rng_k},A{r},{rng_g})").number_format = RP
    hp.cell(row=r, column=8, value=f'=G{r}+G{r}/SUM($G${p0}:$G${pn})*SUMIF({rng_k},"Overhead",{rng_g})').number_format = RP
    hp.cell(row=r, column=9, value=f"=F{r}/H{r}").number_format = TIMES
    for col in range(1, 10):
        hp.cell(row=r, column=col).border = BOX
ht = pn + 1
hp.cell(row=ht, column=2, value="TOTAL").font = BOLD
for col in (6, 7, 8):
    L = get_column_letter(col)
    c = hp.cell(row=ht, column=col, value=f"=SUM({L}{p0}:{L}{pn})")
    c.font, c.number_format = BOLD, RP
hp.cell(row=ht, column=9, value=f"=F{ht}/H{ht}").number_format = TIMES
hp.cell(row=ht + 1, column=2, value="Cek: total = harga diajukan")
hp.cell(row=ht + 1, column=6, value=f'=IF(F{ht}={A["harga"]},"OK","TIDAK SAMA")').font = BOLD
widths(hp, [5, 62, 9, 10, 15, 15, 15, 20, 12])

# ---------------------------------------------------------------- Ringkasan (first tab)
rk = wb.create_sheet("Ringkasan", 0)
rk["A1"] = "RINGKASAN - BI JAMPURNAS 2026 (INTERNAL, bukan untuk klien)"
rk["A1"].font = Font(name="Calibri", bold=True, size=13)
SUMMARY = [
    ("Harga diajukan ke BI", f"={A['harga']}", RP),
    ("Status pajak", f"={A['pkp']}", None),
    ("Diterima setelah pajak (PPh 23)", f"={AFTER_TAX}", RP),
    ("Jumlah biaya", f"={COST}", RP),
    ("PROFIT NETT (setelah pajak dan komisi)", f"={PROFIT}", RP),
    ("Margin nett terhadap harga", f"={MARGIN}", PCT),
    ("Status margin", f"=RAB!$B${refs['r13']}", None),
    ("Harga untuk margin nett 30%", f"=RAB!$B${PRICE_ROWS[0.30]}", RP),
    ("Harga untuk margin nett 35%", f"=RAB!$B${PRICE_ROWS[0.35]}", RP),
    ("Harga untuk margin nett 40%", f"=RAB!$B${PRICE_ROWS[0.40]}", RP),
    ("Cek harga per layanan di proposal", f"='Harga Proposal'!F{ht + 1}", None),
]
for i, (label, f, fmt) in enumerate(SUMMARY, start=3):
    rk.cell(row=i, column=1, value=label).border = BOX
    c = rk.cell(row=i, column=2, value=f)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if label.startswith(("PROFIT", "Margin nett")):
        for col in (1, 2):
            rk.cell(row=i, column=col).font = BOLD
            rk.cell(row=i, column=col).fill = KEY_FILL
notes_row = 3 + len(SUMMARY) + 1
NOTES = [
    "Catatan:",
    "1. Klien: Bank Indonesia Kantor Pusat, acara JAMPURNAS (Jambore Pejuang Rupiah Nasional) 2026.",
    "2. Tim: 1 juri dosen public speaking, 1 juri psikolog, 1 project manager (Galih 28 Okt 2026).",
    "3. Fee juri tetap Rp2,5 juta per orang, sudah termasuk hari acara dan rubrik penilaian 46 karya. Tidak ada fee offline terpisah.",
    "4. PM Rp2 juta dan administrasi Rp400rb ke Coreitera, sesuai kesepakatan.",
    "5. TJ Non-PKP, jadi tidak ada PPN. PPh 23 2% dihitung sebagai biaya.",
    "6. Tab 'Harga Proposal' masih draf dan satu-satunya angka yang boleh keluar ke BI.",
]
for i, n in enumerate(NOTES):
    c = rk.cell(row=notes_row + i, column=1, value=n)
    if i == 0:
        c.font = BOLD
widths(rk, [52, 20])

wb.save(OUT)
print("saved", OUT)
