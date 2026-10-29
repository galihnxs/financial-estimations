"""Builds the BI JAMPURNAS 2026 jury proposal from the CBP Rupiah proposal template.

Usage: python build_proposal.py <template.docx> <out.docx>
Keeps the cover frame, profile, legality, products and attachments of the template,
replaces the program section, and reads service prices from the RAB workbook
("Harga Proposal" tab), so the proposal never shows internal costs.
Run build_rab.py and recalc the workbook in Excel first.
"""
import copy
import sys
from pathlib import Path

import docx
import openpyxl
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
RAB = HERE.parent / "RAB - BI Jampurnas 2026 (Internal).xlsx"
COVER = HERE / "cover_jampurnas.png"
NAVY = RGBColor(0x1C, 0x45, 0x87)
HEAD_HEX, ZEBRA_HEX, STRIPE_HEX, RULE_HEX = "475571", "E4E8F0", "F5F7FB", "D5DBE6"
WHITE, MUTED = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x6B, 0x7A, 0x99)

template, out = sys.argv[1], sys.argv[2]
d = docx.Document(template)
body = d.element.body


def find_par(start):
    for p in d.paragraphs:
        if p.text.strip().startswith(start):
            return p
    raise LookupError(start)


T = {
    "h1": find_par("Pelatihan Metode Belajar Kreatif"),
    "h2": find_par("Latar Belakang Program"),
    "sub": find_par("Tujuan Program"),
    "minor": find_par("Materi 1 - Kesiapan"),
    "body": find_par("Literasi keuangan merupakan"),
    "bullet": find_par("Keterlibatan aktif peserta"),
}
T = {k: copy.deepcopy(v._p) for k, v in T.items()}
anchor = find_par("Lampiran-lampiran Dokumen")._p

# Drop the old program section: from its H1 up to the attachments H1.
start = find_par("Pelatihan Metode Belajar Kreatif")._p
el = start
while el is not anchor:
    nxt = el.getnext()
    body.remove(el)
    el = nxt

# Fix a stray double comma in the template's product text.
for p in d.paragraphs:
    if "penyusunan asesmen, , serta" in p.text:
        for r in p.runs:
            r.text = r.text.replace("penyusunan asesmen, , serta", "penyusunan asesmen, serta")
        if "penyusunan asesmen, , serta" in p.text:  # comma split across runs
            full = p.text.replace("penyusunan asesmen, , serta", "penyusunan asesmen, serta").rstrip(" -")
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = full


def ppr_flag(ppr, tag):
    """Add an on/off paragraph property right after pStyle, where the schema expects it."""
    if ppr.find(qn(tag)) is not None:
        return
    style = ppr.find(qn("w:pStyle"))
    ppr.insert(0 if style is None else list(ppr).index(style) + 1, OxmlElement(tag))


def add(kind, *runs, page_break=False, keep_next=False):
    """Clone a template paragraph. runs: strings, or (text, bold) tuples."""
    p = copy.deepcopy(T[kind])
    rpr = p.find(".//" + qn("w:rPr"))
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    ppr = p.find(qn("w:pPr"))
    if page_break and ppr is not None:
        ppr_flag(ppr, "w:pageBreakBefore")
    if (keep_next or kind in ("h1", "h2", "sub", "minor")) and ppr is not None:
        ppr_flag(ppr, "w:keepNext")  # never strand a heading at the page bottom
    for item in runs:
        text, bold = (item, None) if isinstance(item, str) else item
        r = OxmlElement("w:r")
        if rpr is not None:
            new_rpr = copy.deepcopy(rpr)
            if bold is not None:
                for tag in ("w:b", "w:bCs"):
                    for old in new_rpr.findall(qn(tag)):
                        new_rpr.remove(old)
                    if bold:
                        new_rpr.append(OxmlElement(tag))
            r.append(new_rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p.append(r)
    anchor.addprevious(p)
    return p


def shade(cell, hex_fill):
    xml_child(cell._tc.get_or_add_tcPr(), "w:shd", TCPR_ORDER, val="clear", color="auto", fill=hex_fill)


TBLPR_ORDER = ["w:tblStyle", "w:tblW", "w:jc", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"]
TCPR_ORDER = ["w:tcW", "w:gridSpan", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:vAlign"]


def xml_child(parent, tag, order=None, **attrs):
    """Add (or replace) a child, placed where the Word schema expects it when an order is given."""
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    if order:
        for old in parent.findall(qn(tag)):
            parent.remove(old)
        rank = order.index(tag)
        later = [c for c in parent if c.tag in {qn(t) for t in order[rank + 1:]}]
        if later:
            later[0].addprevious(el)
            return el
    parent.append(el)
    return el


def cell_text(cell, value, *, color, bold=False, size=10.5, align=None):
    """value: a string, or (title, subtitle) for a bold title with a small grey line under it."""
    title, subtitle = (value, None) if isinstance(value, str) else value
    lines = [(title, bold or subtitle is not None, size, color)]
    if subtitle:
        lines.append((subtitle, False, 9, MUTED))
    for k, (text, is_bold, pt, rgb) in enumerate(lines):
        para = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        fmt = para.paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(1 if subtitle and k == 0 else 0), 1.0
        if align is not None:
            para.alignment = align
        run = para.add_run(text)
        run.font.name, run.font.size, run.bold = "Calibri", Pt(pt), is_bold
        run.font.color.rgb = rgb
    return cell


def table(rows, col_widths_cm, footer=(), align_right=(), align_center=()):
    """rows[0] is the header. footer: styles for the last rows ("dark" or "light"); their
    label cells merge across every column except the last. Moves the table before the anchor."""
    tbl = d.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    xml_child(tblPr, "w:tblW", TBLPR_ORDER, w=sum(int(w * 567) for w in col_widths_cm), type="dxa")
    xml_child(tblPr, "w:tblLayout", TBLPR_ORDER, type="fixed")
    # Only horizontal rules: a strong line under the table, hairlines between rows.
    borders = xml_child(tblPr, "w:tblBorders", TBLPR_ORDER)
    for edge, val, sz, color in [("top", "nil", 0, "auto"), ("left", "nil", 0, "auto"),
                                 ("bottom", "single", 12, HEAD_HEX), ("right", "nil", 0, "auto"),
                                 ("insideH", "single", 4, RULE_HEX), ("insideV", "nil", 0, "auto")]:
        attrs = {"val": val} if val == "nil" else {"val": val, "sz": sz, "space": 0, "color": color}
        xml_child(borders, f"w:{edge}", **attrs)
    margins = xml_child(tblPr, "w:tblCellMar", TBLPR_ORDER)
    for edge, w in [("top", 110), ("left", 130), ("bottom", 110), ("right", 130)]:
        xml_child(margins, f"w:{edge}", w=w, type="dxa")
    grid = tbl._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(col_widths_cm[j] * 567)))

    n_body_end = len(rows) - len(footer)
    for i, row in enumerate(rows):
        tr = tbl.rows[i]._tr
        trPr = tr.get_or_add_trPr()
        xml_child(trPr, "w:cantSplit")
        if i == 0:
            xml_child(trPr, "w:tblHeader")
        style = footer[i - n_body_end] if i >= n_body_end else None
        if style:
            label_end = len(row) - 2
            tbl.cell(i, 0).merge(tbl.cell(i, label_end))
        for j, value in enumerate(row):
            if style and 0 < j <= len(row) - 2:
                continue  # merged into the label cell
            cell = tbl.cell(i, j)
            cell.width = docx.shared.Cm(col_widths_cm[j] if not (style and j == 0) else sum(col_widths_cm[:-1]))
            xml_child(cell._tc.get_or_add_tcPr(), "w:vAlign", TCPR_ORDER, val="center")
            align = 2 if j in align_right else 1 if j in align_center else None
            if i == 0:
                shade(cell, HEAD_HEX)
                cell_text(cell, value, color=WHITE, bold=True, size=10, align=align)
            elif style == "dark":
                shade(cell, HEAD_HEX)
                cell_text(cell, value, color=WHITE, bold=True, size=11.5, align=2 if j else None)
            elif style == "light":
                shade(cell, ZEBRA_HEX)
                cell_text(cell, value, color=NAVY, bold=True, align=2 if j else None)
            else:
                if i % 2 == 0:
                    shade(cell, STRIPE_HEX)
                cell_text(cell, value, color=NAVY, align=align)
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next = i < len(rows) - 1  # keep the table on one page
    anchor.addprevious(tbl._tbl)
    spacer = OxmlElement("w:p")
    anchor.addprevious(spacer)
    return tbl


def rupiah(n):
    return "Rp" + f"{round(n):,}".replace(",", ".")


# ---------------------------------------------------------------- prices from the RAB
wb = openpyxl.load_workbook(RAB, data_only=True)
hp = wb["Harga Proposal"]
services = []
for row in hp.iter_rows(min_row=5, values_only=True):
    no, name, vol, unit, price, total = row[:6]
    if not isinstance(no, int):
        break
    if total is None:
        raise SystemExit("RAB has no cached values. Recalculate it in Excel first.")
    services.append((no, name, vol, unit, price, total))
grand = sum(s[5] for s in services)
asumsi = {r[0]: r[1] for r in wb["Asumsi"].iter_rows(min_row=3, values_only=True) if r[0]}
agreed = asumsi["Harga yang diajukan ke BI (total)"]
karya = asumsi["Jumlah karya yang dinilai"]
assert grand == agreed, (grand, agreed)

# ---------------------------------------------------------------- content
add("h1", "Dewan Juri Eksternal Apresiasi CBP Rupiah, JAMPURNAS 2026", page_break=True)

add("h2", "Latar Belakang Program")
add("body", "JAMPURNAS (Jambore Pejuang Rupiah Nasional) adalah ajang tahunan Bank Indonesia untuk memperkuat jejaring, kolaborasi, "
    "dan edukasi Cinta, Bangga, dan Paham (CBP) Rupiah. Salah satu agendanya adalah Apresiasi CBP Rupiah, yaitu penghargaan bagi "
    "satuan kerja kas dan Pejuang Rupiah yang paling berkontribusi dalam meningkatkan literasi Rupiah di masyarakat.")
add("body", "Kategori Edukator CBP Rupiah Terbaik menilai kemampuan Pejuang Rupiah dalam menyampaikan edukasi. Sebanyak 70% nilai "
    "berasal dari video simulasi edukasi, yang mencakup gestur dan ekspresi, intonasi, serta substansi pesan. Penilaian seperti ini "
    "membutuhkan juri eksternal yang memahami dua hal sekaligus: cara seseorang berbicara di depan publik, dan cara pesan tersebut "
    "dipahami serta diterima oleh audiens.")
add("body", "Teman Journey menawarkan dua juri eksternal dengan sudut pandang yang saling melengkapi. Dosen public speaking menilai "
    "cara penyampaian, sedangkan psikolog menilai seberapa mudah pesan dipahami dan seberapa kuat pesan tersebut mendorong audiens "
    "untuk bertindak. Kedua juri menggunakan rubrik yang sama, sehingga setiap peserta dinilai dengan ukuran yang sama.")

add("sub", "Tujuan Program")
for line in [
    "Memberikan penilaian yang objektif, konsisten, dan dapat dipertanggungjawabkan untuk seluruh karya peserta.",
    "Melengkapi penilaian juri internal Bank Indonesia dengan sudut pandang ahli komunikasi publik dan psikologi.",
    "Memastikan pemenang yang terpilih benar-benar mampu menyampaikan edukasi CBP Rupiah secara efektif kepada masyarakat.",
]:
    add("bullet", line)

add("h2", "Ruang Lingkup Penjurian")
add("body", f"Kedua juri menilai seluruh {karya} karya peserta kategori Edukator CBP Rupiah Terbaik, lalu hadir sebagai juri pada hari "
    "acara JAMPURNAS. Rubrik penilaian disusun dari kriteria dalam Kerangka Acuan Kerja (ToR) Apresiasi CBP Rupiah 2026.")
table([
    ["Kriteria dalam ToR", "Bobot", "Fokus penilaian juri eksternal"],
    ["Portofolio", "20%", "Mengikuti ketentuan Bank Indonesia"],
    [("Simulasi edukasi: gestur dan ekspresi", "Bahasa tubuh, kontak mata, ekspresi wajah"), "20%", "Dosen public speaking"],
    [("Simulasi edukasi: intonasi", "Volume, tempo, penekanan kata"), "10%", "Dosen public speaking"],
    [("Simulasi edukasi: substansi", "Bahasa sesuai audiens, alur materi, kalimat punchline"), "40%", "Dosen public speaking dan psikolog"],
    [("Evaluasi dan tindak lanjut", "Self-assessment dan rencana perbaikan"), "10%", "Psikolog"],
], [7.6, 1.8, 7.1], align_center=(1,))
add("body", "Bobot akhir dan pembagian penilaian antara juri internal dan juri eksternal mengikuti ketentuan Bank Indonesia.")

add("h2", "Komponen Layanan")
add("body", "Program terdiri atas tiga layanan.")

add("minor", "Layanan 1 - Juri Dosen Public Speaking")
add("body", f"Dosen public speaking menilai {karya} video simulasi edukasi dengan rubrik dan form penilaian yang telah disetujui "
    "Bank Indonesia. Fokus penilaiannya adalah cara peserta menyampaikan materi:", keep_next=True)
for line in [
    "Gestur, ekspresi, dan kontak mata dengan audiens.",
    "Intonasi, tempo, dan kejelasan suara.",
    "Alur penyampaian yang efektif dan tidak bertele-tele, serta kekuatan kalimat punchline.",
]:
    add("bullet", line, keep_next=True)
add("body", "Dosen juga hadir sebagai juri pada hari acara JAMPURNAS.")

add("minor", "Layanan 2 - Juri Psikolog")
add("body", f"Psikolog Teman Journey menilai {karya} karya yang sama dengan rubrik yang sama. Fokus penilaiannya adalah bagaimana "
    "pesan diterima oleh audiens:")
for line in [
    "Pemilihan bahasa yang sesuai dengan audiens dan mudah dipahami.",
    "Kekuatan pesan dalam mempersuasi dan mendorong audiens untuk bertindak.",
    "Kualitas refleksi diri peserta dalam self-assessment dan rencana tindak lanjut.",
]:
    add("bullet", line)
add("body", "Psikolog juga hadir sebagai juri pada hari acara JAMPURNAS.")

add("minor", "Layanan 3 - Manajemen Program dan Pendampingan Acara")
add("body", "Seorang project manager Teman Journey menjadi narahubung tunggal bagi panitia Bank Indonesia. Tugasnya meliputi:")
for line in [
    "Menyiapkan rubrik dan form penilaian, lalu meminta persetujuan Bank Indonesia.",
    "Mengatur jadwal penilaian dan kehadiran kedua juri.",
    "Merekap nilai kedua juri dan menyerahkannya kepada panitia tepat waktu.",
    "Mendampingi juri pada hari acara.",
]:
    add("bullet", line)

add("h2", "Alur Pelaksanaan")
table([
    ["Tahap", "Kegiatan", "Waktu"],
    ["1", "Penandatanganan SPK dan briefing bersama panitia Bank Indonesia", "Minggu ke-1"],
    ["2", "Penyusunan rubrik dan form penilaian, lalu persetujuan Bank Indonesia", "Minggu ke-1"],
    ["3", f"Penerimaan {karya} karya peserta dari panitia", "Sesuai jadwal panitia"],
    ["4", "Penilaian karya oleh kedua juri", "Sebelum hari acara"],
    ["5", "Penyerahan rekap nilai kepada panitia", "Sebelum hari acara, sesuai jadwal panitia"],
    ["6", "Penjurian pada hari acara JAMPURNAS", "Hari acara"],
], [1.6, 9.2, 5.7], align_center=(0,))
add("body", "Tanggal pasti disepakati bersama Bank Indonesia setelah SPK ditandatangani.")

add("h2", "Hasil yang Diterima Bank Indonesia")
for line in [
    "Rubrik dan form penilaian yang disusun dari kriteria ToR dan disetujui Bank Indonesia.",
    f"Nilai {karya} karya peserta dari dua juri eksternal.",
    "Rekap nilai seluruh peserta yang siap digabungkan dengan nilai juri internal.",
    "Kehadiran dua juri eksternal pada hari acara JAMPURNAS.",
]:
    add("bullet", line)

add("h2", "Tim Juri")
add("body", "Juri dosen public speaking adalah dosen yang mengajar dan berpraktik di bidang komunikasi publik. Nama dan profilnya "
    "dikirim kepada Bank Indonesia sebelum pelaksanaan. Juri psikolog ditetapkan dari tim psikolog Teman Journey yang tersertifikasi "
    "HIMPSI berikut.", keep_next=True)
table([
    ["Nama", "Profesi", "S2", "S1"],
    ["Riskia Ramadhina Sukriananda", "Psikolog", "Universitas Indonesia", "Universitas Indonesia"],
    ["Amanda Nurshadrina", "Psikolog", "Universitas Indonesia", "Universitas Padjadjaran"],
    ["Laras Sita Persitia", "Psikolog Klinis Dewasa", "Universitas Indonesia", "Universitas Brawijaya"],
    ["Fasya Khalida Suffah", "Psikolog Klinis Dewasa", "Universitas Indonesia", "Universitas Indonesia"],
    ["Nadira Khairunnisa", "Psikolog Klinis Dewasa", "Universitas Indonesia", "Universitas Indonesia"],
    ["Ninette Putri Mustika", "Psikolog", "Universitas Indonesia", "Universitas Indonesia"],
    ["Rhajiv Nur Ilham", "Psikolog Klinis Dewasa", "Universitas Indonesia", "Universitas Indonesia"],
], [5.0, 3.8, 3.8, 3.9])

add("h2", "Independensi dan Kerahasiaan")
add("body", "Kedua juri tidak memiliki hubungan kerja dengan satuan kerja kas yang dinilai dan menilai setiap karya secara independen. "
    "Karya, data peserta, dan nilai hanya digunakan untuk keperluan Apresiasi CBP Rupiah 2026 dan hanya diserahkan kepada Bank "
    "Indonesia. Psikolog Teman Journey bekerja sesuai Kode Etik Psikologi Indonesia (HIMPSI), dan data pribadi peserta dikelola sesuai "
    "UU No. 27 Tahun 2022 tentang Pelindungan Data Pribadi.")

add("h2", "Proposal Anggaran", page_break=True)
add("body", "Rincian harga untuk setiap layanan adalah sebagai berikut.")
# One line under each service name so BI sees what the price buys.
WHAT = {
    1: f"Penilaian {karya} karya dengan rubrik, kehadiran di hari acara",
    2: f"Penilaian {karya} karya dengan rubrik, kehadiran di hari acara",
    3: "Rubrik dan form penilaian, koordinasi juri, rekap nilai, pendampingan acara",
}
TITLE = {1: "Juri dosen public speaking", 2: "Juri psikolog", 3: "Manajemen program dan pendampingan acara"}
rows = [["No", "Layanan", "Volume", "Harga Satuan", "Jumlah"]]
for no, name, vol, unit, price, total in services:
    rows.append([str(no), (TITLE[no], WHAT[no]), f"{vol} {unit}", rupiah(price), rupiah(total)])
rows.append(["Total*", "", "", "", rupiah(grand)])
table(rows, [0.9, 7.2, 2.3, 2.8, 3.3], footer=("dark",), align_right=(3, 4), align_center=(0, 2))
add("body", "*Harga sudah termasuk pajak yang berlaku. Biaya perjalanan dan akomodasi juri, apabila acara diselenggarakan di luar "
    "Jabodetabek, ditanggung oleh panitia. Ketentuan pembayaran mengikuti SPK. Penawaran berlaku 30 hari sejak tanggal proposal.")

# Attachments heading starts on a fresh page.
apPr = anchor.find(qn("w:pPr"))
if apPr is not None:
    ppr_flag(apPr, "w:pageBreakBefore")

# ---------------------------------------------------------------- cover image
swapped = False
for section in d.sections:
    for hdr in (section.first_page_header, section.header):
        for rel in hdr.part.rels.values():
            if rel.reltype.endswith("/image") and rel.target_ref.endswith("image7.png"):
                rel.target_part._blob = COVER.read_bytes()
                swapped = True
assert swapped, "cover image not found"

# The template's TOC lists styles by name, which Word on Mac does not match. Use outline levels.
for instr in body.iter(qn("w:instrText")):
    if instr.text and instr.text.strip().startswith("TOC"):
        instr.text = ' TOC \\o "1-2" \\h \\z \\u '

d.core_properties.title = "Proposal Dewan Juri Eksternal Apresiasi CBP Rupiah JAMPURNAS 2026 - Bank Indonesia"
d.save(out)
print("saved", out)
